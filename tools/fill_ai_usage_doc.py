from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("docs/4-AI工具使用说明（填写版-橘源通）（2026年版）.docx")


ROWS = [
    [
        "1",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年3月17日-4月25日",
        "立项构思与需求梳理：明确柑橘种植、分级、预警、溯源的 MVP 功能边界。",
        "请围绕柑橘产业数字化，梳理一个微信小程序+后端+AI工程的核心功能和最小可运行版本。",
        "给出批次管理、农事日志、病害识别、果实分级定价、风险预警、二维码溯源等模块建议。",
        "团队结合赛题方向和现有联调能力，删除过大的电商/IoT设想，保留可实现闭环。",
        "采纳约35%构思内容；不涉及代码。",
    ],
    [
        "2",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年3月20日-4月21日",
        "系统架构与接口文档：辅助整理微信小程序、NestJS 后端、Python AI 工程的分层职责。",
        "请根据微信小程序、NestJS、Python AI 的技术栈，整理模块划分、接口返回约定和联调注意事项。",
        "建议统一 {code,message,data} 返回，区分 /api 业务接口与 /ai 推理接口，并保留兼容路由。",
        "接口路径、字段和兼容关系由人工按代码逐项核对，文档表述按实际实现修正。",
        "采纳约30%文档结构；不涉及代码。",
    ],
    [
        "3",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年3月22日-4月18日",
        "病害识别开发：辅助生成 PyTorch 训练/预测脚手架、标签映射和推理流程参考。",
        "请设计柑橘病害图像分类的训练、评估和预测流程，输出 Python 模块组织和关键函数建议。",
        "给出数据集加载、模型工厂、预测输出 label/confidence/severity/advice 的代码草案。",
        "人工重构数据路径、标签兼容、模型加载、异常兜底和业务字段，核心识别结果由传统模型产生。",
        "AI生成代码直接采纳约18%；最终代码占比不超过25%。",
    ],
    [
        "4",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年3月25日-4月20日",
        "果实分级与定价：辅助设计颜色、缺陷、大小、成熟度等特征和规则定价说明。",
        "请为柑橘果实分级定价设计可解释特征、等级规则、价格区间计算和返回字段。",
        "提出 colorScore、defectRatio、sizeScore、maturityScore、grade、price range、factors 等字段。",
        "人工确定阈值、价格系数、字段兼容和 fallback，定价由规则/传统模型计算，LLM只负责解释。",
        "AI生成代码直接采纳约15%；最终代码占比不超过25%。",
    ],
    [
        "5",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年3月28日-4月19日",
        "风险预警规则：辅助梳理天气、农事日志、病害记录对风险等级的影响。",
        "请根据天气、病害记录、农事日志和批次阶段，设计农业风险评估规则与可解释输出。",
        "给出 level、levelText、reason、suggestion、规则命中说明等输出建议。",
        "团队人工实现规则引擎和阈值，补充兼容接口；AI文本只用于解释，不替代风险打级。",
        "AI生成代码直接采纳约12%；最终代码占比不超过25%。",
    ],
    [
        "6",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年4月1日-4月23日",
        "LLM解释与建议：辅助设计提示词模板、结构化输出和失败回退策略。",
        "请为病害解释、果实分级解释、风险说明、农事建议和问答设计中文、保守、可执行的提示词。",
        "建议统一 prompts、结构化 JSON 输出、低置信度提示、失败时 deterministic fallback。",
        "人工加入“仅基于输入字段、不得臆造、农业建议保守表达”的约束，并做密钥脱敏和错误处理。",
        "AI生成代码直接采纳约20%；最终代码占比不超过25%。",
    ],
    [
        "7",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年4月5日-4月24日",
        "小程序页面与请求封装：辅助生成页面状态、表单校验和 API 调用逻辑参考。",
        "请根据接口清单生成微信小程序页面调用流程，覆盖病害识别、果实分级、风险、溯源和农事日志。",
        "给出 services 请求封装、页面 loading/error 状态、结果展示字段和 smoke test 建议。",
        "团队按微信开发者工具实际调试，调整页面路径、字段兼容、上传流程和交互细节。",
        "AI生成代码直接采纳约22%；最终代码占比不超过25%。",
    ],
    [
        "8",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年4月10日-4月24日",
        "测试与验收：辅助整理构建命令、接口 smoke test 清单和风险点。",
        "请根据当前接口与模块，整理后端构建、AI脚本、关键接口字段和失败回退的验收清单。",
        "输出 npm run build、AI脚本 --help、关键接口字段、LLM失败回退等检查项。",
        "人工执行或核对可运行命令，按实际项目状态删除未落地项，保留 MVP 稳定性检查。",
        "采纳约30%测试清单文字；不涉及代码。",
    ],
    [
        "9",
        "OpenAI ChatGPT/Codex（GPT-5），客户端/API 访问，2026年4月20日-4月25日",
        "参赛材料润色：辅助整理作品简介、创新点、AI使用边界和说明材料。",
        "请将“橘源通”的项目功能、AI分层架构和竞赛材料说明润色为正式、简洁、可审查的中文表述。",
        "给出作品简介、创新描述、AI工具使用说明和“传统模型/规则负责核心判定”的表述。",
        "人工核对事实与代码现状，删除夸大宣传，补充“AI写作/代码采纳比例不超过25%”说明。",
        "采纳约40%文字表达；不涉及新增代码。",
    ],
]


APPENDIX = [
    "总体声明：本作品中AI工具主要用于需求梳理、代码草案参考、文档润色和解释性文本生成；经人工筛选、重构、测试和事实核对后，最终直接采纳的AI生成代码估算约20%-22%，未超过25%。病害分类、果实评分/定价和风险打级均由传统模型或规则完成，LLM不替代核心确定性计算。",
    "序号1的佐证材料：对应 docs/需求分析文档.md、docs/橘源通-概要设计文档（项目版）.md、docs/项目全量总结与接口现状.md。AI输出为模块候选清单，最终功能边界由团队按MVP可运行目标人工筛选。",
    "序号2的佐证材料：对应 docs/api-interface-list.md、docs/ai-solution.md、docs/ai-backend-integration.md。AI辅助整理架构表达，接口路径和字段以实际代码为准人工核对。",
    "序号3的佐证材料：对应 ai/src/disease/、ai/scripts/train_disease.py、ai/scripts/predict_disease.py、backend/src/modules/disease。AI草案仅作脚手架参考，核心模型训练、标签兼容、推理兜底由人工实现和调试。",
    "序号4的佐证材料：对应 ai/src/fruit_grade/、ai/config/fruit_grade_rules.example.yaml、backend/src/modules/price。AI建议用于字段和说明草案，等级阈值、价格规则、兼容字段由人工确定。",
    "序号5的佐证材料：对应 ai/src/risk/、backend/src/modules/risk、docs/risk-engine-design.md。风险等级由规则引擎确定，AI仅辅助梳理解释文本和文档描述。",
    "序号6的佐证材料：对应 backend/src/llm、ai/src/llm、ai/prompts/、docs/ai-solution.md。提示词集中管理，要求仅基于输入字段作答，失败时使用确定性回退。",
    "序号7的佐证材料：对应 pages/、services/、app.json、docs/测试与验收手册.md。AI辅助生成页面状态和请求流程建议，最终页面路径、字段映射、上传联调由人工调试。",
    "序号8的佐证材料：对应 docs/测试与验收手册.md、AGENTS.md 中验证命令，以及 backend/package.json、ai/scripts/。AI辅助整理验收清单，命令和风险结论由人工核对。",
    "序号9的佐证材料：对应 docs/2-作品信息概要表（必填模板）（2026年版）.docx、docs/project-status.md、docs/ai-innovation-analysis.md。AI辅助文字润色，作品名称、功能边界和AI职责边界由团队确认。",
]


def set_font(run, size=10.5, bold=False, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text="", size=10.5, bold=False, align=None, after=4):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)
    return paragraph


def set_cell_text(cell, text, size=7.0, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(str(text).split("\n")):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, value=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name in ["top", "start", "bottom", "end"]:
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def make_fixed_table(table, widths_twips):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    existing_grid = table._tbl.tblGrid
    if existing_grid is not None:
        table._tbl.remove(existing_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_twips[idx])
            set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)

    add_para(doc, "中国大学生计算机设计大赛", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_para(doc, "AI工具使用说明（2026年版）", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(
        doc,
        "作品编号：待填写（以报名系统/学校汇总编号为准）        作品名称：橘源通：面向柑橘产业的智能识别、分级定价与可信溯源小程序",
        size=10.5,
        after=8,
    )

    add_para(doc, "一、AI工具使用情况", size=12, bold=True, after=5)
    field_names = [
        "序号",
        "AI工具的名称、版本、访问方式（网页、API或客户端），使用时间",
        "使用AI工具的环节与目的",
        "关键提示词",
        "AI回复的关键内容",
        "AI回复的人工修改说明",
        "采纳比例与说明",
    ]
    for row_data in ROWS:
        add_para(doc, f"序号{row_data[0]}", size=10.5, bold=True, after=2)
        for field, value in zip(field_names[1:], row_data[1:]):
            add_para(doc, f"{field}：{value}", size=9.2, after=2)
        add_para(doc, "", size=4, after=2)

    doc.add_page_break()
    add_para(doc, "附录1：作品文件夹示例", size=12, bold=True, after=4)
    for text in [
        "2026012345-参赛总文件夹",
        "├── 2026012345-01作品与答辩材料",
        "├── 2026012345-02素材与源码",
        "├── 2026012345-03设计与开发文档",
        "└── 2026012345-04作品演示视频",
    ]:
        add_para(doc, text, size=10.5, after=1)

    add_para(doc, "附录2：AI工具使用佐证材料与采纳声明", size=12, bold=True, after=4)
    for text in APPENDIX:
        add_para(doc, text, size=8.8, after=2)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
